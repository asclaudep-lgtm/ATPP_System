"""PDO Dispatcher v2 — Kanban board with rich cards, timeline, documents."""

from datetime import datetime, date
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QScrollArea, QFrame, QGroupBox, QMessageBox, QInputDialog,
    QSplitter, QTextEdit, QSizePolicy, QGridLayout,
    QLineEdit, QFormLayout, QTableWidget, QTableWidgetItem, QDialog,
)
from PyQt6.QtCore import Qt, pyqtSignal, QTimer, QSize
from PyQt6.QtGui import QFont, QColor, QPalette

from modules import pdo_module
from database.models import PDOStatus

STATUS_COLORS = {
    PDOStatus.NEW:                '#6c757d',
    PDOStatus.OMTS_REVIEW:        '#16a085',
    PDOStatus.TECH_DEPT:          '#2980b9',
    PDOStatus.FEASIBLE:           '#27ae60',
    PDOStatus.NOT_FEASIBLE:       '#c0392b',
    PDOStatus.DEPUTY_APPROVAL:    '#8e44ad',
    PDOStatus.APPROVED:           '#27ae60',
    PDOStatus.WITH_TECHNOLOGIST:  '#e67e22',
    PDOStatus.MTP_SIGNED:         '#2980b9',
    PDOStatus.READY_FOR_SHOP:     '#27ae60',
    PDOStatus.IN_SHOP:            '#8e44ad',
    PDOStatus.QC:                 '#e74c3c',
    PDOStatus.CLOSED:             '#2c3e50',
    PDOStatus.CANCELLED:          '#c0392b',
}

KANBAN_COLUMNS = [
    PDOStatus.NEW,
    PDOStatus.OMTS_REVIEW,
    PDOStatus.TECH_DEPT,
    PDOStatus.DEPUTY_APPROVAL,
    PDOStatus.APPROVED,
    PDOStatus.IN_SHOP,
    PDOStatus.QC,
]

PRIORITY_STARS = {1: '🔴🔴🔴', 2: '🟠🟠', 3: '🟡', 4: '⚪', 5: '—'}


class PDOOrderCard(QFrame):
    """Rich order card for Kanban column."""

    clicked = pyqtSignal(int)

    def __init__(self, order_data, parent=None):
        super().__init__(parent)
        self.order_id = order_data['id']
        self.setFrameShape(QFrame.Shape.StyledPanel)
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setToolTip(
            f"№{order_data['number']} | {order_data['product']}\n"
            f"Статус: {order_data['status']} | Срок: {order_data.get('due_date', '—')}")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(8, 6, 8, 6)
        layout.setSpacing(3)

        # Status strip
        status_color = STATUS_COLORS.get(
            getattr(PDOStatus, order_data['status'].upper().replace(' ', '_'),
                    None) if order_data['status'] in [
                        s.value for s in PDOStatus] else None,
            '#999')

        # Top row: number + priority
        top = QHBoxLayout()
        num = QLabel(order_data['number'])
        num.setFont(QFont('Arial', 10, QFont.Weight.Bold))
        top.addWidget(num)
        top.addStretch()
        prio = order_data.get('priority', 3)
        stars = PRIORITY_STARS.get(prio, '')
        if stars:
            prio_lbl = QLabel(stars)
            prio_lbl.setFont(QFont('Arial', 7))
            top.addWidget(prio_lbl)
        layout.addLayout(top)

        # Product + aircraft
        prod = QLabel(order_data.get('product', '—'))
        prod.setWordWrap(True)
        prod.setFont(QFont('Arial', 9))
        layout.addWidget(prod)

        ac = order_data.get('aircraft_type', '')
        if ac:
            ac_lbl = QLabel(f'✈ {ac[:20]}')
            ac_lbl.setStyleSheet('color: #1976d2; font-size: 9px;')
            layout.addWidget(ac_lbl)

        # Nomenclature summary
        items_total = order_data.get('items_total', 0)
        if items_total:
            items_row = QHBoxLayout()
            items_row.addWidget(QLabel(
                f'📋 {items_total} поз.'))
            n_feas = order_data.get('items_feasible', 0)
            n_not = order_data.get('items_not_feasible', 0)
            if n_feas:
                lbl = QLabel(f' +{n_feas}')
                lbl.setStyleSheet('color: #27ae60; font-size: 10px;')
                items_row.addWidget(lbl)
            if n_not:
                lbl = QLabel(f' -{n_not}')
                lbl.setStyleSheet('color: #c0392b; font-weight: bold; font-size: 10px;')
                items_row.addWidget(lbl)
            items_row.addStretch()
            layout.addLayout(items_row)

        # Bottom: qty + due date
        bot = QHBoxLayout()
        qty_lbl = QLabel(f"×{order_data.get('qty', 1)} шт")
        qty_lbl.setStyleSheet('color: #555; font-size: 11px;')
        bot.addWidget(qty_lbl)
        bot.addStretch()

        due_str = order_data.get('due_date', '') or ''
        due_lbl = QLabel(due_str)
        due_lbl.setStyleSheet('font-size: 10px;')
        if due_str:
            try:
                d = date.fromisoformat(due_str)
                today = date.today()
                if d < today:
                    due_lbl.setStyleSheet(
                        'color: #c0392b; font-weight: bold; font-size: 10px;')
                elif d == today:
                    due_lbl.setStyleSheet(
                        'color: #e67e22; font-weight: bold; font-size: 10px;')
            except ValueError:
                pass
        bot.addWidget(due_lbl)
        layout.addLayout(bot)

        # Memos line
        memo_count = order_data.get('memo_count', 0)
        omts = order_data.get('omts_memo', '')
        deputy = order_data.get('deputy_memo', '')
        if omts or deputy:
            memo_row = QHBoxLayout()
            if omts:
                lbl = QLabel(f'📄 У:{omts[:12]}')
                lbl.setStyleSheet('font-size: 9px; color: #16a085;')
                memo_row.addWidget(lbl)
            if deputy:
                lbl = QLabel(f'П:{deputy[:12]}')
                lbl.setStyleSheet('font-size: 9px; color: #8e44ad;')
                memo_row.addWidget(lbl)
            memo_row.addStretch()
            layout.addLayout(memo_row)

        # Color indicator on left side
        self.setStyleSheet(
            f'PDOOrderCard {{ border-left: 4px solid {status_color}; '
            f'background: #fff; margin: 2px 4px; border-radius: 4px; }}'
            f'PDOOrderCard:hover {{ background: #f0f4ff; }}')

    def mousePressEvent(self, ev):
        self.clicked.emit(self.order_id)
        super().mousePressEvent(ev)


class PDOKanbanColumn(QGroupBox):
    """Scrollable column of order cards for one status."""

    def __init__(self, status: PDOStatus, parent=None):
        super().__init__(parent)
        self.status = status
        color = STATUS_COLORS.get(status, '#999')
        title = f'{status.value}  '
        self.setTitle(title)
        self.setStyleSheet(
            f'PDOKanbanColumn {{ '
            f'border-top: 3px solid {color}; '
            f'background: #f8f9fa; border-radius: 6px; '
            f'padding-top: 12px; }}')

        layout = QVBoxLayout(self)
        layout.setContentsMargins(2, 4, 2, 2)

        self._scroll = QScrollArea()
        self._scroll.setWidgetResizable(True)
        self._scroll.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        self._scroll.setStyleSheet(
            'QScrollArea { border: none; background: transparent; }')

        self._container = QWidget()
        self._card_layout = QVBoxLayout(self._container)
        self._card_layout.setSpacing(2)
        self._card_layout.setContentsMargins(0, 0, 0, 0)
        self._card_layout.addStretch()
        self._scroll.setWidget(self._container)
        layout.addWidget(self._scroll)

    def clear_cards(self):
        while self._card_layout.count() > 1:
            item = self._card_layout.takeAt(0)
            if item.widget():
                item.widget().deleteLater()

    def add_card(self, order_data) -> PDOOrderCard:
        card = PDOOrderCard(order_data)
        self._card_layout.insertWidget(
            self._card_layout.count() - 1, card)
        return card


class PDODispatcherWidget(QWidget):
    """Kanban dispatcher with rich cards, detail panel, and document generation."""

    def __init__(self, db_manager, user, parent=None):
        super().__init__(parent)
        self.db_manager = db_manager
        self.user = user
        self._init_ui()
        self.refresh()

        self._refresh_timer = QTimer()
        self._refresh_timer.timeout.connect(self.refresh)
        self._refresh_timer.start(30000)

    def _init_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(4, 4, 4, 4)
        layout.setSpacing(6)

        # Header
        hdr = QHBoxLayout()
        title = QLabel('📋 Диспетчер ПДО — Производственные заказы')
        title.setFont(QFont('Arial', 14, QFont.Weight.Bold))
        hdr.addWidget(title)
        hdr.addStretch()

        new_btn = QPushButton('+ Новый заказ')
        new_btn.setStyleSheet(
            'QPushButton { background: #27ae60; color: white; '
            'border: none; padding: 8px 16px; border-radius: 4px; '
            'font-weight: bold; }')
        new_btn.clicked.connect(self._create_order)
        hdr.addWidget(new_btn)

        refresh_btn = QPushButton('↻ Обновить')
        refresh_btn.clicked.connect(self.refresh)
        hdr.addWidget(refresh_btn)

        self._stats_label = QLabel('')
        self._stats_label.setStyleSheet('color: #666; margin-left: 12px;')
        hdr.addWidget(self._stats_label)
        layout.addLayout(hdr)

        # Main splitter: Kanban | Detail panel
        splitter = QSplitter(Qt.Orientation.Horizontal)

        # Kanban scroll area
        kanban_scroll = QScrollArea()
        kanban_scroll.setWidgetResizable(True)
        kanban_widget = QWidget()
        self._kanban_layout = QHBoxLayout(kanban_widget)
        self._kanban_layout.setSpacing(6)
        self._kanban_layout.setContentsMargins(0, 0, 0, 0)

        self._columns = {}
        for st in KANBAN_COLUMNS:
            col = PDOKanbanColumn(st)
            self._kanban_layout.addWidget(col)
            self._columns[st] = col

        kanban_scroll.setWidget(kanban_widget)
        splitter.addWidget(kanban_scroll)

        # Detail panel
        detail = self._build_detail_panel()
        splitter.addWidget(detail)

        splitter.setSizes([800, 380])
        layout.addWidget(splitter, stretch=1)

    def _build_detail_panel(self):
        panel = QWidget()
        panel.setStyleSheet(
            'background: #fff; border-radius: 6px; padding: 8px;')
        layout = QVBoxLayout(panel)
        layout.setSpacing(8)

        # Title
        self._detail_title = QLabel('Выберите заказ в канбан-доске')
        self._detail_title.setFont(QFont('Arial', 12, QFont.Weight.Bold))
        self._detail_title.setWordWrap(True)
        layout.addWidget(self._detail_title)

        # Info grid
        self._info_grid = QGridLayout()
        self._info_labels = {}
        fields = [
            ('status', 'Статус'), ('product', 'Изделие'),
            ('qty', 'Количество'), ('due_date', 'Срок'),
            ('priority', 'Приоритет'), ('customer', 'Заказчик'),
            ('technologist', 'Технолог'),
            ('mtp_signed', 'МТП'), ('shop', 'Цех'),
            ('qty_done', 'Выполнено'), ('qty_scrap', 'Брак'),
        ]
        for i, (key, label) in enumerate(fields):
            lbl = QLabel(f'{label}:')
            lbl.setStyleSheet('color: #666; font-size: 11px;')
            val = QLabel('—')
            val.setWordWrap(True)
            self._info_labels[key] = val
            self._info_grid.addWidget(lbl, i // 2, (i % 2) * 2)
            self._info_grid.addWidget(val, i // 2, (i % 2) * 2 + 1)
        layout.addLayout(self._info_grid)

        # Action buttons
        self._action_layout = QHBoxLayout()

        self._sign_mtp_btn = QPushButton('✍ Подписать МТП')
        self._sign_mtp_btn.setStyleSheet(
            'QPushButton { background: #2980b9; color: white; '
            'border: none; padding: 6px 14px; border-radius: 4px; }')
        self._sign_mtp_btn.clicked.connect(self._sign_mtp)
        self._action_layout.addWidget(self._sign_mtp_btn)

        self._release_btn = QPushButton('📦 Передать в цех')
        self._release_btn.setStyleSheet(
            'QPushButton { background: #8e44ad; color: white; '
            'border: none; padding: 6px 14px; border-radius: 4px; }')
        self._release_btn.clicked.connect(self._release_to_shop)
        self._action_layout.addWidget(self._release_btn)

        self._close_btn = QPushButton('✓ Закрыть заказ')
        self._close_btn.clicked.connect(self._close_order)
        self._action_layout.addWidget(self._close_btn)

        self._doc_btn = QPushButton('📄 Печать карты заказа')
        self._doc_btn.clicked.connect(self._print_order_card)
        self._action_layout.addWidget(self._doc_btn)

        self._action_layout.addStretch()
        layout.addLayout(self._action_layout)

        # Hide all action buttons initially
        for b in [self._sign_mtp_btn, self._release_btn,
                   self._close_btn, self._doc_btn]:
            b.hide()

        # Handoff timeline
        timeline_lbl = QLabel('История передач:')
        timeline_lbl.setFont(QFont('Arial', 10, QFont.Weight.Bold))
        layout.addWidget(timeline_lbl)

        self._timeline = QTextEdit()
        self._timeline.setReadOnly(True)
        self._timeline.setMaximumHeight(140)
        self._timeline.setStyleSheet(
            'font-size: 11px; background: #f8f9fa; border-radius: 4px;')
        layout.addWidget(self._timeline)

        return panel

    # ── Refresh ───────────────────────────────────────────────────

    def refresh(self):
        with self.db_manager.get_session() as s:
            total = 0
            for st, col in self._columns.items():
                orders = pdo_module.list_orders_by_status(s, st, limit=50)
                col.clear_cards()
                for o in orders:
                    detail = pdo_module.get_order_detail(s, o.id)
                    # Count feasible/not items
                    items = detail.get('nomenclature', [])
                    feasible_count = sum(
                        1 for i in items if i.get('tech_feasible') is True)
                    not_feasible_count = sum(
                        1 for i in items if i.get('tech_feasible') is False)
                    memos = detail.get('memos', [])
                    card = col.add_card({
                        'id': o.id, 'number': o.number,
                        'product': o.product.designation if o.product else '—',
                        'qty': o.qty,
                        'due_date': str(o.due_date) if o.due_date else '',
                        'priority': o.priority or 3,
                        'status': o.status.value,
                        'aircraft_type': detail.get('aircraft_type', ''),
                        'technologist': detail.get('technologist', ''),
                        'mtp_signed': detail.get('mtp_signed', False),
                        'omts_memo': detail.get('omts_memo_no', ''),
                        'deputy_memo': detail.get('deputy_memo_no', ''),
                        'items_total': len(items),
                        'items_feasible': feasible_count,
                        'items_not_feasible': not_feasible_count,
                        'memo_count': len(memos),
                    })
                    card.clicked.connect(self._show_detail)
                    total += 1
                col.setTitle(f'{st.value}  ({len(orders)})')
            self._stats_label.setText(
                f'Всего заказов: {total}  |  '
                f'Обновлено: {datetime.now().strftime("%H:%M:%S")}')

    # ── Detail ────────────────────────────────────────────────────

    def _show_detail(self, order_id):
        with self.db_manager.get_session() as s:
            detail = pdo_module.get_order_detail(s, order_id)
            if not detail:
                return
            self._current_detail = detail

            self._detail_title.setText(
                f'{detail["number"]} — {detail["product"]}')
            self._info_labels['status'].setText(detail['status'])
            ac = detail.get('aircraft_type', '')
            self._info_labels['product'].setText(
                f'{detail.get("product_name", "")}'
                + (f'  ✈ {ac}' if ac else ''))
            self._info_labels['qty'].setText(
                f'{detail["qty"]} шт')
            self._info_labels['due_date'].setText(
                detail.get('due_date', '—') or '—')
            self._info_labels['priority'].setText(
                PRIORITY_STARS.get(detail.get('priority', 3), '—'))
            self._info_labels['customer'].setText(
                detail.get('customer', '—') or '—')
            self._info_labels['technologist'].setText(
                detail.get('technologist', '—'))
            mtp_status = '✅ Подписан' if detail.get('mtp_signed') else '❌ Не подписан'
            self._info_labels['mtp_signed'].setText(mtp_status)
            self._info_labels['shop'].setText(
                detail.get('released_to_shop', '—') or '—')
            self._info_labels['qty_done'].setText(
                f'{detail.get("qty_done", 0)} шт')
            self._info_labels['qty_scrap'].setText(
                f'{detail.get("qty_scrap", 0)} шт')

            # Nomenclature table
            items = detail.get('nomenclature', [])
            nomen_lines = ['=== НОМЕНКЛАТУРА ===']
            for item in items:
                feas = {True: '✅', False: '❌', None: '⬜'}.get(
                    item.get('tech_feasible'), '?')
                kd = '📄' if item.get('kd_ready') else '  '
                mat = item.get('material_name', '')[:30]
                nomen_lines.append(
                    f'{feas} {kd} {item["designation"][:25]} — '
                    f'{item.get("name","")[:20]} ×{item["qty"]}шт'
                    + (f'  [{mat}]' if mat else ''))

            # Memos
            memos = detail.get('memos', [])
            if memos:
                nomen_lines.append('')
                nomen_lines.append('=== СЛУЖЕБНЫЕ ЗАПИСКИ ===')
                for m in memos:
                    t = 'УКАЗАНИЕ' if m['memo_type'] == 'У' else 'ПРИКАЗ'
                    nomen_lines.append(
                        f'📄 {t} {m["memo_number"]} — {m["from_dept"]}'
                        + (f': {m["content"][:50]}' if m.get('content') else ''))

            # Handoffs
            nomen_lines.append('')
            nomen_lines.append('=== ПЕРЕДАЧИ ===')
            for h in detail.get('handoffs', []):
                nomen_lines.append(
                    f'{h["from"]} → {h["to"]}  [{h["status"]}]')

            self._timeline.setPlainText('\n'.join(nomen_lines))

            # Show/hide action buttons
            st = detail['status']
            role = self.user.get('role', '')
            self._sign_mtp_btn.setVisible(
                role in ('technologist', 'admin')
                and st in ('Новый', 'У технолога'))
            self._release_btn.setVisible(
                role in ('admin', 'technologist', 'engineer')
                and st in ('МТП подписан', 'Готов к передаче в цех'))
            self._close_btn.setVisible(
                role in ('qc', 'admin') and st == 'ОТК')
            self._doc_btn.setVisible(True)

    # ── Actions ───────────────────────────────────────────────────

    def _create_order(self):
        from PyQt6.QtWidgets import QDialog, QFormLayout, QTableWidget
        from database.models import Product

        dlg = QDialog(self)
        dlg.setWindowTitle('Новый заказ ПДО')
        dlg.setMinimumSize(650, 500)
        layout = QVBoxLayout(dlg)

        form = QFormLayout()
        des_edit = QLineEdit()
        des_edit.setPlaceholderText('Обозначение основного изделия')
        form.addRow('Изделие *:', des_edit)

        ac_edit = QLineEdit()
        ac_edit.setPlaceholderText('Ту-204-300 / Ил-76МД / ...')
        form.addRow('Тип ВС:', ac_edit)

        customer_edit = QLineEdit()
        customer_edit.setPlaceholderText('Заказчик')
        form.addRow('Заказчик:', customer_edit)

        ws_edit = QLineEdit()
        ws_edit.setPlaceholderText('Краткое описание работ')
        form.addRow('Объём работ:', ws_edit)

        layout.addLayout(form)

        QLabel('Номенклатура (позиции к изготовлению):').setStyleSheet(
            'font-weight: bold; margin-top: 8px;')
        layout.addWidget(QLabel('Номенклатура (позиции к изготовлению):'))

        tbl = QTableWidget(0, 3)
        tbl.setHorizontalHeaderLabels(['Обозначение', 'Наименование', 'Кол-во'])
        tbl.setColumnWidth(0, 200)
        tbl.setColumnWidth(1, 200)
        tbl.setColumnWidth(2, 60)
        layout.addWidget(tbl)

        btn_row = QHBoxLayout()
        add_btn = QPushButton('+ Добавить позицию')
        add_btn.clicked.connect(lambda: tbl.insertRow(tbl.rowCount()))
        btn_row.addWidget(add_btn)
        del_btn = QPushButton('− Удалить')
        del_btn.clicked.connect(lambda: tbl.removeRow(tbl.currentRow()))
        btn_row.addWidget(del_btn)
        btn_row.addStretch()
        layout.addLayout(btn_row)

        # Pre-populate with one row
        tbl.insertRow(0)
        tbl.setItem(0, 0, QTableWidgetItem(des_edit.text()))

        ok_btn = QPushButton('Создать заказ')
        ok_btn.setStyleSheet(
            'QPushButton { background: #27ae60; color: white; '
            'border: none; padding: 8px; border-radius: 4px; font-weight: bold; }')
        layout.addWidget(ok_btn)

        def on_create():
            des = des_edit.text().strip()
            if not des:
                QMessageBox.warning(dlg, 'Ошибка', 'Введите обозначение изделия.')
                return
            with self.db_manager.get_session() as s:
                p = s.query(Product).filter(
                    Product.designation == des,
                    Product.is_deleted == False).first()
                if p is None:
                    QMessageBox.warning(dlg, 'Ошибка', f'Изделие "{des}" не найдено.')
                    return

                nomenclature = []
                for r in range(tbl.rowCount()):
                    d = tbl.item(r, 0)
                    n = tbl.item(r, 1)
                    q = tbl.item(r, 2)
                    if d and d.text().strip():
                        try:
                            qty_val = int(q.text()) if q and q.text().strip() else 1
                        except ValueError:
                            qty_val = 1
                        nomenclature.append({
                            'designation': d.text().strip(),
                            'name': n.text().strip() if n else '',
                            'qty': qty_val,
                        })

                total_qty = sum(ni['qty'] for ni in nomenclature) or 1
                order = pdo_module.create_order(
                    s, product_id=p.id, qty=total_qty,
                    due_date=date.today(),
                    customer=customer_edit.text().strip(),
                    aircraft_type=ac_edit.text().strip(),
                    work_scope=ws_edit.text().strip(),
                    created_by=self.user.get('id', 1),
                    nomenclature=nomenclature,
                )
            dlg.accept()
            QMessageBox.information(
                self, 'Заказ создан',
                f'{order.number}\n'
                f'Изделие: {p.designation} — {p.name}\n'
                f'Позиций: {len(nomenclature)}\n'
                f'Тип ВС: {ac_edit.text().strip() or "—"}\n'
                f'Статус: {order.status.value}')
            self.refresh()

        ok_btn.clicked.connect(on_create)
        dlg.exec()

    def _sign_mtp(self):
        if not hasattr(self, '_current_detail'):
            return
        oid = self._current_detail['id']
        tp_id_str, ok = QInputDialog.getText(
            self, 'Подпись МТП', 'ID техпроцесса:')
        if not ok or not tp_id_str.strip():
            return
        comment, ok2 = QInputDialog.getText(
            self, 'Комментарий', 'Примечание к подписи:')
        if not ok2:
            return
        try:
            with self.db_manager.get_session() as s:
                pdo_module.sign_mtp(
                    s, order_id=oid,
                    tech_process_id=int(tp_id_str.strip()),
                    signed_by=self.user.get('id', 1),
                    comment=comment,
                )
            self.refresh()
            self._show_detail(oid)
        except Exception as e:
            QMessageBox.critical(self, 'Ошибка', str(e))

    def _release_to_shop(self):
        if not hasattr(self, '_current_detail'):
            return
        oid = self._current_detail['id']
        shop, ok = QInputDialog.getText(
            self, 'Передача в цех',
            'Номер цеха / участка:\n\n'
            'При передаче будут сформированы:\n'
            '• Маршрутная карта (МК)\n'
            '• Операционные карты (ОК)\n'
            '• МТП с подписью технолога\n'
            '• Ведомость оснастки (ВО)\n'
            '• Ведомость материалов (ВМ)',
            text='Цех 1')
        if not ok or not shop.strip():
            return
        with self.db_manager.get_session() as s:
            pdo_module.release_to_shop(
                s, order_id=oid, shop=shop.strip(),
                released_by=self.user.get('id', 1),
            )
        self.refresh()
        self._show_detail(oid)
        QMessageBox.information(
            self, 'Передано',
            f'Заказ передан в {shop}.\n'
            f'Документы сформированы.\n'
            f'Уведомление отправлено мастеру цеха.')

    def _close_order(self):
        if not hasattr(self, '_current_detail'):
            return
        oid = self._current_detail['id']
        reply = QMessageBox.question(
            self, 'Закрытие заказа',
            f'Заказ {self._current_detail["number"]}\n'
            f'Выполнено: {self._current_detail.get("qty", 0)} шт\n\n'
            f'Подтвердите закрытие.',
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if reply != QMessageBox.StandardButton.Yes:
            return
        with self.db_manager.get_session() as s:
            pdo_module.close_order(
                s, order_id=oid,
                qty_done=self._current_detail.get('qty', 0),
            )
        self.refresh()

    def _print_order_card(self):
        from modules.doc_generator import DocumentGenerator
        if not hasattr(self, '_current_detail'):
            return
        oid = self._current_detail['id']

        with self.db_manager.get_session() as s:
            order = s.query(
                __import__('database.models', fromlist=['ProductionOrder'])
                .ProductionOrder
            ).get(oid)
            if order is None:
                return

            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side
            from config import EXPORT_DIR

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f'Заказ {order.number}'
            ws.column_dimensions['A'].width = 22
            ws.column_dimensions['B'].width = 40

            ws['A1'] = f'КАРТА ЗАКАЗА {order.number}'
            ws['A1'].font = Font(size=14, bold=True)
            ws.merge_cells('A1:B1')

            fields = [
                ('Номер заказа', order.number),
                ('Изделие', (order.product.designation + ' — ' +
                             order.product.name)
                 if order.product else '—'),
                ('Количество', f'{order.qty} шт'),
                ('Срок', str(order.due_date) if order.due_date else '—'),
                ('Приоритет', str(order.priority)),
                ('Заказчик', order.customer or '—'),
                ('Статус', order.status.value),
                ('Технолог', (order.technologist.full_name
                              if order.technologist else 'Не назначен')),
                ('МТП', 'Подписан' if order.mtp_signed_at else 'Не подписан'),
                ('Цех', order.released_to_shop or '—'),
                ('Выполнено', f'{order.qty_done} шт'),
                ('Брак', f'{order.qty_scrap} шт'),
            ]
            for i, (label, value) in enumerate(fields, 3):
                ws[f'A{i}'] = label
                ws[f'A{i}'].font = Font(bold=True)
                ws[f'B{i}'] = value

            path = EXPORT_DIR / f'order_card_{order.number}.xlsx'
            path.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(path))

        QMessageBox.information(self, 'Готово', f'Карта заказа сохранена:\n{path}')

    def closeEvent(self, ev):
        self._refresh_timer.stop()
        super().closeEvent(ev)

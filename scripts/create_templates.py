"""Создать шаблоны .docx для генерации документов ТП.

Генерирует два шаблона в resources/templates/ktd_docx/:
  1. title_page_template.docx  — титульный лист
  2. route_card_template.docx  — маршрутная карта (МК)
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

from docx import Document
from docx.shared import Cm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUTPUT = Path(__file__).parent.parent / 'resources' / 'templates' / 'ktd_docx'


def _set_cell_border(cell, **kwargs):
    """Set cell border properties."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        from lxml import etree
        tcBorders = etree.SubElement(tcPr, qn('w:tcBorders'))
    for edge, val in kwargs.items():
        element = tcBorders.find(qn(f'w:{edge}'))
        if element is None:
            from lxml import etree
            element = etree.SubElement(tcBorders, qn(f'w:{edge}'))
        element.set(qn('w:val'), val.get('val', 'single'))
        element.set(qn('w:sz'), val.get('sz', '4'))
        element.set(qn('w:color'), val.get('color', '000000'))


def create_title_page():
    """Титульный лист ГОСТ."""
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.0)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)

    # Empty spacer paragraphs for top alignment
    for _ in range(6):
        doc.add_paragraph('')

    # Company name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('{{company}}')
    run.font.size = Pt(14)
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('АТПП — Автоматизация технологической подготовки производства')
    run.font.size = Pt(10)

    doc.add_paragraph('')
    doc.add_paragraph('')

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('ТЕХНОЛОГИЧЕСКИЙ ПРОЦЕСС')
    run.font.size = Pt(16)
    run.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('{{tp_number}}')
    run.font.size = Pt(14)

    doc.add_paragraph('')

    # Product info table
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    fields = [
        ('Обозначение изделия:', '{{designation}}'),
        ('Наименование изделия:', '{{name}}'),
        ('Материал:', '{{material}}'),
        ('Вид технологии:', '{{technology_type}}'),
        ('Дата разработки:', '{{date}}'),
    ]

    for i, (label, placeholder) in enumerate(fields):
        cell0 = table.cell(i, 0)
        cell0.width = Cm(5)
        cell0.text = label
        for p in cell0.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(11)

        cell1 = table.cell(i, 1)
        cell1.width = Cm(10)
        cell1.text = placeholder
        for p in cell1.paragraphs:
            p.runs[0].font.size = Pt(11)

    doc.add_paragraph('')

    # Signatures block
    table2 = doc.add_table(rows=4, cols=4)
    table2.style = 'Table Grid'
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ['Должность', 'Фамилия', 'Подпись', 'Дата']
    for j, h in enumerate(headers):
        cell = table2.cell(0, j)
        cell.text = h
        for p in cell.paragraphs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(10)

    roles = [('Разработал', '{{author}}'), ('Проверил', '{{checker}}'), ('Утвердил', '{{approver}}')]
    for i, (role, person) in enumerate(roles):
        table2.cell(i + 1, 0).text = role
        table2.cell(i + 1, 1).text = person
        table2.cell(i + 1, 2).text = ''
        table2.cell(i + 1, 3).text = ''

    path = OUTPUT / 'title_page_template.docx'
    OUTPUT.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    print(f'Created: {path}')
    return path


def create_route_card():
    """Маршрутная карта (МК) ГОСТ 3.1118-82."""
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(0.5)

    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(2)

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('МАРШРУТНАЯ КАРТА')
    run.font.size = Pt(14)
    run.font.bold = True

    # GOST reference
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('по ГОСТ 3.1118-82')
    run.font.size = Pt(9)
    run.font.italic = True

    # Info block
    info_table = doc.add_table(rows=1, cols=6)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    info_cells = info_table.rows[0].cells
    labels = ['Обозначение', 'Наименование', 'Материал', 'Твёрдость', 'Масса', 'Номер ТП']
    values = ['{{designation}}', '{{name}}', '{{material}}', '{{hardness}}', '{{mass}}', '{{tp_number}}']

    for j, (label, value) in enumerate(zip(labels, values)):
        p = info_cells[j].paragraphs[0]
        run = p.add_run(label + '\n')
        run.font.size = Pt(7)
        run.font.bold = True
        run = p.add_run(value)
        run.font.size = Pt(9)

    doc.add_paragraph('')

    # Operations table
    table = doc.add_table(rows=1, cols=10)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ['№ оп.', 'Наименование операции', 'Оборудование', 'Профессия',
               'Разряд', 'Тшт', 'Тпз', 'То', 'Тв', 'Цех']
    header_cells = table.rows[0].cells
    for j, h in enumerate(headers):
        p = header_cells[j].paragraphs[0]
        run = p.add_run(h)
        run.font.size = Pt(8)
        run.font.bold = True

    # Placeholder row template marker
    row = table.add_row()
    for j, placeholder in enumerate([
        '{{op_number}}', '{{op_name}}', '{{op_equipment}}', '{{op_profession}}',
        '{{op_grade}}', '{{op_t_piece}}', '{{op_t_setup}}',
        '{{op_t_main}}', '{{op_t_aux}}', '{{op_shop}}'
    ]):
        p = row.cells[j].paragraphs[0]
        run = p.add_run(placeholder)
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(150, 150, 150)

    # Add a few more placeholder rows
    for _ in range(3):
        table.add_row()

    doc.add_paragraph('')

    # Totals
    p = doc.add_paragraph()
    run = p.add_run('Суммарное Тшт: {{total_t_piece}} мин    |    Суммарное Тпз: {{total_t_setup}} мин')
    run.font.size = Pt(10)
    run.font.bold = True

    # Signatures block
    doc.add_paragraph('')
    sig_table = doc.add_table(rows=2, cols=5)
    sig_table.style = 'Table Grid'
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    sig_headers = ['', 'Фамилия', 'Подпись', 'Дата', '']
    for j, h in enumerate(sig_headers):
        p = sig_table.rows[0].cells[j].paragraphs[0]
        run = p.add_run(h)
        run.font.size = Pt(9)
        run.font.bold = True

    sig_roles = [('Разработал:', '{{author}}'), ('Проверил:', '{{checker}}'), ('Нормоконтроль:', '{{normer}}'), ('Утвердил:', '{{approver}}')]
    for i, (role, person) in enumerate(sig_roles):
        if i < 4:
            sig_table.rows[1].cells[i].text = f'{role}\n{person}'

    path = OUTPUT / 'route_card_template.docx'
    doc.save(str(path))
    print(f'Created: {path}')
    return path


if __name__ == '__main__':
    create_title_page()
    create_route_card()
    print('\nDone. Templates created in:', OUTPUT)

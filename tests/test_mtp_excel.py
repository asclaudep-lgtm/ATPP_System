"""Тесты генератора МТП в .xlsx по форме УЗГА."""
import io

import pytest
from sqlalchemy import event
from sqlalchemy.engine import Engine

from database.db_manager import DatabaseManager
from database.models import (
    Operation, Product, TechProcess, TPStatus, User,
)
from modules.mtp_excel import (
    _filter_collapse_intermediate,
    generate_mtp_excel,
)
from modules.production import release_to_production


@event.listens_for(Engine, 'connect')
def _enable_sqlite_fk(dbapi_connection, _):  # pragma: no cover
    try:
        cur = dbapi_connection.cursor()
        cur.execute('PRAGMA foreign_keys=ON;')
        cur.close()
    except Exception:
        pass


@pytest.fixture
def session(tmp_path):
    db_path = tmp_path / 'test_mtp.db'
    dm = DatabaseManager(f'sqlite:///{db_path}')
    dm.init_database()
    with dm.get_session() as s:
        # Базовые сущности
        admin = s.query(User).filter_by(username='admin').first()
        product = Product(
            designation='51-XX.00.0001',
            name='Деталь тестовая',
            blank_dimensions='100x50x5 мм',
        )
        s.add(product)
        s.flush()
        tp = TechProcess(
            number=f'{product.designation} ТП',
            product_id=product.id,
            status=TPStatus.APPROVED,
            version='1.0',
            author_id=admin.id,
        )
        s.add(tp)
        s.flush()

        # 5 операций: разные типы
        for i, (num, name) in enumerate([
            ('005', 'Заготовительная'),
            ('010', 'Контрольная'),
            ('015', 'Слесарная'),
            ('020', 'Промывочная'),
            ('025', 'Контрольная'),
        ]):
            s.add(Operation(
                tech_process_id=tp.id,
                number=num,
                name=name,
                sort_order=i,
                include_in_mtp=True,
            ))
        s.commit()

        admin_dict = {
            'id': admin.id, 'username': admin.username, 'role': admin.role
        }
        s.info['admin'] = admin_dict
        s.info['tp_id'] = tp.id
        yield s


def test_filter_collapse_intermediate_keeps_last_check():
    class O:
        def __init__(self, name):
            self.name = name
            self.include_in_mtp = True

    ops = [
        O('Заготовительная'),
        O('Контрольная'),
        O('Слесарная'),
        O('Промывочная'),
        O('Контрольная'),
    ]
    out = _filter_collapse_intermediate(ops)
    names = [o.name for o in out]
    assert names == ['Заготовительная', 'Слесарная', 'Контрольная']


def test_filter_collapse_handles_empty_list():
    assert _filter_collapse_intermediate([]) == []


def test_generate_mtp_excel_with_wo(session, tmp_path):
    """Генератор должен собрать xlsx с встроенным штрих-кодом."""
    tp_id = session.info['tp_id']
    admin = session.info['admin']
    wo = release_to_production(
        session, user=admin, tech_process_id=tp_id, qty_total=10,
        customer_order='Заказ-Т-001',
    )
    session.flush()

    out_path = tmp_path / 'mtp.xlsx'
    data = generate_mtp_excel(
        session,
        tech_process_id=tp_id,
        work_order_id=wo.id,
        out_path=out_path,
        collapse_intermediate=True,
    )
    assert len(data) > 1000
    assert out_path.exists()
    assert out_path.stat().st_size == len(data)


def test_generate_mtp_excel_without_wo(session, tmp_path):
    """Без наряда — формируется черновик без штрих-кода."""
    tp_id = session.info['tp_id']
    out_path = tmp_path / 'mtp_draft.xlsx'
    data = generate_mtp_excel(
        session, tech_process_id=tp_id, out_path=out_path,
    )
    assert len(data) > 1000
    assert out_path.exists()


def test_generate_mtp_excel_full_and_compact_differ(session, tmp_path):
    """Свёрнутая и полная версии должны отличаться по размеру/содержимому."""
    tp_id = session.info['tp_id']
    admin = session.info['admin']
    wo = release_to_production(
        session, user=admin, tech_process_id=tp_id, qty_total=10,
    )
    session.flush()

    full = generate_mtp_excel(
        session, tech_process_id=tp_id, work_order_id=wo.id,
        collapse_intermediate=False,
    )
    compact = generate_mtp_excel(
        session, tech_process_id=tp_id, work_order_id=wo.id,
        collapse_intermediate=True,
    )
    # В обоих случаях получаем валидный xlsx (PK-сигнатура zip)
    assert full[:2] == b'PK'
    assert compact[:2] == b'PK'

    # Проверим количество операций через openpyxl
    import openpyxl
    wb_full = openpyxl.load_workbook(io.BytesIO(full))
    wb_compact = openpyxl.load_workbook(io.BytesIO(compact))
    ws_full = wb_full.active
    ws_compact = wb_compact.active

    def _count_ops(ws):
        n = 0
        for r in range(10, ws.max_row + 1):
            v = ws.cell(row=r, column=4).value
            if v and str(v).strip().isdigit():
                n += 1
        return n

    assert _count_ops(ws_full) == 5     # все 5 операций
    assert _count_ops(ws_compact) == 3  # без 010 «Контрольная» и 020 «Промывочная»


def test_generate_mtp_excel_unknown_tp_raises(session):
    from modules.mtp_excel import MTPExcelError
    with pytest.raises(MTPExcelError):
        generate_mtp_excel(session, tech_process_id=999999)

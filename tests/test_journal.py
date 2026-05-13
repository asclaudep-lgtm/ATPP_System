"""Тесты единого журнала регистрации ТП/МТП."""
from database.models import (
    User, Product, TechProcess, RegistrationJournal, TPStatus, TPType,
)


def _seed_product_and_tp(db):
    with db.get_session() as s:
        u = s.query(User).filter_by(username='admin').first()
        p = Product(designation='ABC.001', name='Тест-деталь')
        s.add(p); s.flush()
        tp = TechProcess(number='TP-001', version='1', author_id=u.id,
                         product_id=p.id, status=TPStatus.DRAFT,
                         tp_type=TPType.SINGLE)
        s.add(tp); s.flush()
        return p.id, tp.id


def test_next_number_starts_at_first(db_manager):
    from modules.journal import next_number
    with db_manager.get_session() as s:
        n1 = next_number(s)
    assert n1 == 'УЗГА.02101.00001'


def test_register_product_assigns_number(db_manager):
    from modules import journal as J
    pid, _ = _seed_product_and_tp(db_manager)
    with db_manager.get_session() as s:
        p = s.query(Product).get(pid)
        e = J.register_product(s, p, user_id=1, executor='Иванов И.И.')
        assert e.tp_number == 'УЗГА.02101.00001'
        assert e.mtp_number == 'УЗГА.02101.00001'
        assert e.entry_no == 1
        assert e.in_tp_journal is True
        assert e.in_mtp_journal is True


def test_register_two_products_get_different_numbers(db_manager):
    from modules import journal as J
    with db_manager.get_session() as s:
        p1 = Product(designation='A.001', name='A')
        p2 = Product(designation='B.001', name='B')
        s.add(p1); s.add(p2); s.flush()
        e1 = J.register_product(s, p1, user_id=1)
        e2 = J.register_product(s, p2, user_id=1)
        n1, n2 = e1.tp_number, e2.tp_number
    assert n1 == 'УЗГА.02101.00001'
    assert n2 == 'УЗГА.02101.00002'


def test_exclude_and_restore(db_manager):
    from modules import journal as J
    pid, _ = _seed_product_and_tp(db_manager)
    with db_manager.get_session() as s:
        p = s.query(Product).get(pid)
        e = J.register_product(s, p, user_id=1)
        eid = e.id
    with db_manager.get_session() as s:
        ok = J.exclude_entry(s, eid, user_id=1, reason='тест')
        assert ok is True
    with db_manager.get_session() as s:
        e = s.query(RegistrationJournal).get(eid)
        assert e.excluded is True
        assert e.excluded_reason == 'тест'
        # active list не должен включать
        active = J.list_entries(s, only_active=True)
        assert eid not in [x.id for x in active]
    with db_manager.get_session() as s:
        ok = J.restore_entry(s, eid, user_id=1)
        assert ok is True
    with db_manager.get_session() as s:
        e = s.query(RegistrationJournal).get(eid)
        assert e.excluded is False
        active = J.list_entries(s, only_active=True)
        assert eid in [x.id for x in active]


def test_export_layouts(db_manager, tmp_path):
    from modules import journal as J
    pid, tpid = _seed_product_and_tp(db_manager)
    with db_manager.get_session() as s:
        p = s.query(Product).get(pid)
        J.register_product(s, p, user_id=1, executor='X')
        rows = J.list_entries(s)
    for layout in ('tp', 'mtp', 'unified'):
        out = tmp_path / f'{layout}.xlsx'
        with db_manager.get_session() as s:
            rows = J.list_entries(s)
            path = J.export_to_excel(s, rows, layout=layout, out_path=out)
        assert path.exists()
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert wb.active.max_row >= 2

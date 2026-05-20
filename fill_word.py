import subprocess
import sys

sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# --- Get unique passport numbers ---
ps_script = r'''
$target = (New-Object -ComObject WScript.Shell).CreateShortcut('C:\Users\InjTeh\Desktop\инженерный центр.lnk').TargetPath
Get-ChildItem $target -Recurse -ErrorAction SilentlyContinue | Where-Object { $_.Name -match '^53-74\.80\.\d+\.\d+' } | ForEach-Object {
    if ($_.Name -match '^(53-74\.80\.\d+\.\d+)') {
        Write-Output $Matches[1]
    }
}
'''

result = subprocess.run(
    ['powershell', '-NoProfile', '-Command', ps_script],
    capture_output=True, text=True, timeout=120
)
lines = result.stdout.strip().split('\n')
unique = sorted(set(line.strip() for line in lines if line.strip().startswith('53-74.80.')))
print(f'Unique passports: {len(unique)}')

# --- Open existing document ---
doc_path = r'C:\Users\InjTeh\Desktop\Перечень паспортов 53-74-80.docx'
doc = Document(doc_path)

# Clear old table and paragraphs
for t in doc.tables:
    t._element.getparent().remove(t._element)
for p in doc.paragraphs:
    p._element.getparent().remove(p._element)

# --- Add header ---
title = doc.add_paragraph('Перечень паспортов группы 53-74.80 (64050)')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size = Pt(14)
title.runs[0].font.bold = True

count_p = doc.add_paragraph(f'Всего: {len(unique)} паспортов')
count_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

sep = doc.add_paragraph('─' * 60)
sep.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- Build table (4 columns) ---
rows_needed = (len(unique) + 3) // 4
table = doc.add_table(rows=rows_needed + 1, cols=4)
table.style = 'Table Grid'

# Header row
for c_idx, label in enumerate(['№', 'Паспорт', '№', 'Паспорт']):
    cell = table.rows[0].cells[c_idx]
    cell.text = label
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.size = Pt(9)

# Fill data
for i, num in enumerate(unique):
    row = (i % rows_needed) + 1
    col = (i // rows_needed) * 2
    # Number column
    cell_num = table.rows[row].cells[col]
    cell_num.text = str(i + 1)
    for run in cell_num.paragraphs[0].runs:
        run.font.size = Pt(8)
    # Passport column
    cell_pas = table.rows[row].cells[col + 1]
    cell_pas.text = num
    for run in cell_pas.paragraphs[0].runs:
        run.font.size = Pt(8)

doc.save(doc_path)
print(f'Saved: {doc_path}')
print('Printing...')

# Print
subprocess.run(['powershell', '-Command',
    "Start-Process -FilePath '" + doc_path + "' -Verb Print"], timeout=30)
print('Sent to printer.')
